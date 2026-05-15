from io import BytesIO

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

from app.models.enums import ConversionMode
from app.pipeline.layout.reading_order import sort_blocks_reading_order
from app.pipeline.models import DocumentModel, ImageBlock, TableBlock, TableCellStyle, TextBlock, TextSpan
from app.word.layout_math import bbox_width_inches, pt_to_docx_indent, pt_to_space_before
from app.word.oxml_utils import set_cell_shading, set_paragraph_shading
from app.word.styles import ALIGNMENT_MAP, HEADING_STYLES


def build_docx(doc_model: DocumentModel, mode: ConversionMode | None = None) -> bytes:
    if getattr(doc_model, "use_visual_grid", False):
        from app.word.grid_renderer import render_document_grid

        return render_document_grid(doc_model)

    document = Document()
    preserve_layout = mode is None or mode == ConversionMode.ACCURATE

    for idx, page in enumerate(doc_model.pages):
        if idx > 0:
            document.add_page_break()

        blocks = sort_blocks_reading_order(page.blocks)
        for block in blocks:
            if isinstance(block, TextBlock):
                if block.is_header or block.is_footer:
                    continue
                _add_text_block(document, block, page.width, preserve_layout)
            elif isinstance(block, TableBlock):
                _add_table_block(document, block)
            elif isinstance(block, ImageBlock):
                _add_image_block(document, block, page.width)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _hex_to_rgb(hex_color: str | None) -> RGBColor | None:
    if not hex_color:
        return None
    h = hex_color.lstrip("#")
    if len(h) != 6:
        return None
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _apply_run_style(run, span: TextSpan, block: TextBlock) -> None:
    size = span.font_size or block.font_size
    if size:
        run.font.size = Pt(min(size, 72))
    if span.bold or block.bold:
        run.bold = True
    if span.italic or block.italic:
        run.italic = True
    if span.underline or block.underline:
        run.underline = True
    rgb = _hex_to_rgb(span.color or block.color)
    if rgb:
        run.font.color.rgb = rgb


def _add_text_block(
    document: Document,
    block: TextBlock,
    page_width: float,
    preserve_layout: bool,
) -> None:
    if block.type == "heading":
        p = document.add_paragraph()
        p.style = HEADING_STYLES.get(block.level, "Heading 1")
    elif block.type == "list_item":
        p = document.add_paragraph(style="List Bullet")
    else:
        p = document.add_paragraph()

    spans = block.spans if block.spans else [
        TextSpan(
            text=block.text,
            font_size=block.font_size,
            color=block.color,
            bold=block.bold,
            italic=block.italic,
            underline=block.underline,
        )
    ]

    if not spans or (len(spans) == 1 and not spans[0].text):
        if block.text:
            run = p.add_run(block.text)
            _apply_run_style(run, spans[0] if spans else TextSpan(text=block.text), block)
    else:
        for span in spans:
            if not span.text:
                continue
            run = p.add_run(span.text)
            _apply_run_style(run, span, block)

    align = ALIGNMENT_MAP.get(block.alignment)
    if align is not None:
        p.alignment = align

    if preserve_layout:
        pf = p.paragraph_format
        pf.left_indent = pt_to_docx_indent(block.left_indent_pt, page_width)
        pf.space_before = pt_to_space_before(block.space_before_pt)
        line_height_pt = max(block.bbox.y1 - block.bbox.y0, block.font_size * 1.1)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line_height_pt)

    if block.background_color:
        set_paragraph_shading(p, block.background_color)


def _add_table_block(document: Document, block: TableBlock) -> None:
    if not block.rows:
        return
    rows = len(block.rows)
    cols = max(len(r) for r in block.rows)
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"

    for r_idx, row in enumerate(block.rows):
        for c_idx, cell_text in enumerate(row):
            if c_idx >= cols:
                continue
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = str(cell_text)
            style: TableCellStyle | None = None
            if block.cell_styles and r_idx < len(block.cell_styles) and c_idx < len(block.cell_styles[r_idx]):
                style = block.cell_styles[r_idx][c_idx]
            if style:
                if style.background_color:
                    set_cell_shading(cell, style.background_color)
                if cell.paragraphs and cell.paragraphs[0].runs:
                    run = cell.paragraphs[0].runs[0]
                    if style.bold:
                        run.bold = True
                    rgb = _hex_to_rgb(style.text_color)
                    if rgb:
                        run.font.color.rgb = rgb
    document.add_paragraph()


def _add_image_block(document: Document, block: ImageBlock, page_width: float) -> None:
    try:
        stream = BytesIO(block.data)
        width_pt = block.bbox.x1 - block.bbox.x0
        width_in = bbox_width_inches(width_pt, page_width) if width_pt > 0 else min(block.width / 96, 6.5)
        document.add_picture(stream, width=Inches(max(0.5, width_in)))
        document.add_paragraph()
    except Exception:
        pass
