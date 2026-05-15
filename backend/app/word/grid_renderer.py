"""Stage 5 DOCX output — invisible layout grid (mirror-like visual reconstruction)."""

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

from app.pipeline.canvas_models import GridPlacement, PageGrid
from app.pipeline.models import DocumentModel, ImageBlock, PageLayout, TableBlock, TableCellStyle, TextBlock, TextSpan
from app.word.layout_math import page_content_width_in
from app.word.oxml_utils import set_cell_margins_zero, set_cell_shading, set_paragraph_shading, set_row_height, set_table_borders_none
from app.word.styles import ALIGNMENT_MAP, HEADING_STYLES


def render_document_grid(doc_model: DocumentModel) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    for idx, page in enumerate(doc_model.pages):
        if idx > 0:
            document.add_page_break()
        if page.grid and page.blocks:
            _render_page_grid(document, page)
        else:
            from app.word.builder import _add_image_block, _add_table_block, _add_text_block

            for block in page.blocks:
                if isinstance(block, TextBlock) and not block.is_header and not block.is_footer:
                    _add_text_block(document, block, page.width, True)
                elif isinstance(block, TableBlock):
                    _add_table_block(document, block)
                elif isinstance(block, ImageBlock):
                    _add_image_block(document, block, page.width)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_page_grid(document: Document, page: PageLayout) -> None:
    grid: PageGrid = page.grid  # type: ignore[assignment]
    scale = page_content_width_in(page.width) / (page.width / 72.0) if page.width else 1.0

    table = document.add_table(rows=grid.row_count, cols=grid.col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders_none(table)
    table.autofit = False

    for ci, w_pt in enumerate(grid.col_widths_pt):
        if ci < len(table.columns):
            table.columns[ci].width = Pt(w_pt * scale)

    for ri, h_pt in enumerate(grid.row_heights_pt):
        if ri < len(table.rows):
            set_row_height(table.rows[ri], h_pt * scale)

    merged: set[tuple[int, int]] = set()

    for placement in grid.placements:
        if placement.is_spacer:
            continue
        if placement.block_index < 0 or placement.block_index >= len(page.blocks):
            continue

        r0, c0 = placement.row, placement.col
        r1 = min(r0 + placement.rowspan - 1, grid.row_count - 1)
        c1 = min(c0 + placement.colspan - 1, grid.col_count - 1)

        if (r0, c0) in merged:
            continue

        try:
            cell = table.cell(r0, c0)
            if placement.rowspan > 1 or placement.colspan > 1:
                end = table.cell(r1, c1)
                cell = cell.merge(end)
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    merged.add((r, c))
        except Exception:
            cell = table.cell(r0, c0)

        set_cell_margins_zero(cell)
        cell.text = ""
        block = page.blocks[placement.block_index]
        _render_block_in_cell(cell, block, page.width, scale)

    document.add_paragraph()


def _hex_to_rgb(hex_color: str | None) -> RGBColor | None:
    if not hex_color:
        return None
    h = hex_color.lstrip("#")
    if len(h) != 6:
        return None
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _apply_run(run, span: TextSpan, block: TextBlock) -> None:
    size = span.font_size or block.font_size
    if size:
        run.font.size = Pt(min(size, 72))
    if span.bold or block.bold:
        run.bold = True
    if span.italic or block.italic:
        run.italic = True
    if span.underline or block.underline:
        run.underline = True
    rgb = _hex_to_rgb(span.color or block.color)
    if rgb:
        run.font.color.rgb = rgb
    if span.font_name or block.font_name:
        run.font.name = (span.font_name or block.font_name or "")[:64]


def _clear_paragraph(p) -> None:
    for run in list(p.runs):
        run._element.getparent().remove(run._element)


def _render_text_in_cell(cell, block: TextBlock) -> None:
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _clear_paragraph(p)

    if block.type == "heading":
        p.style = HEADING_STYLES.get(block.level, "Heading 1")

    spans = block.spans or [
        TextSpan(
            text=block.text,
            font_size=block.font_size,
            color=block.color,
            bold=block.bold,
            italic=block.italic,
            underline=block.underline,
            font_name=block.font_name,
        )
    ]
    for span in spans:
        if not span.text:
            continue
        run = p.add_run(span.text)
        _apply_run(run, span, block)

    align = ALIGNMENT_MAP.get(block.alignment)
    if align is not None:
        p.alignment = align

    line_h = max(block.bbox.y1 - block.bbox.y0, block.font_size * 1.1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(line_h)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    if block.background_color:
        set_paragraph_shading(p, block.background_color)


def _render_table_in_cell(cell, block: TableBlock, scale: float) -> None:
    if not block.rows:
        return
    ncols = max(len(r) for r in block.rows)
    nested = cell.add_table(rows=len(block.rows), cols=ncols)
    nested.style = "Table Grid"
    for r_idx, row in enumerate(block.rows):
        for c_idx, text in enumerate(row):
            if c_idx >= ncols:
                continue
            nc = nested.rows[r_idx].cells[c_idx]
            nc.text = str(text)
            style = None
            if block.cell_styles and r_idx < len(block.cell_styles) and c_idx < len(block.cell_styles[r_idx]):
                style = block.cell_styles[r_idx][c_idx]
            if style and style.background_color:
                set_cell_shading(nc, style.background_color)
            if style and nc.paragraphs and nc.paragraphs[0].runs:
                run = nc.paragraphs[0].runs[0]
                if style.bold:
                    run.bold = True
                rgb = _hex_to_rgb(style.text_color)
                if rgb:
                    run.font.color.rgb = rgb


def _render_image_in_cell(cell, block: ImageBlock, page_width: float, scale: float) -> None:
    try:
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        run = p.add_run()
        stream = BytesIO(block.data)
        w_pt = block.bbox.x1 - block.bbox.x0
        width_in = (w_pt / 72.0) * scale if w_pt > 0 else min(block.width / 96, 6.5)
        run.add_picture(stream, width=Inches(max(0.25, width_in)))
    except Exception:
        pass


def _render_block_in_cell(cell, block, page_width: float, scale: float) -> None:
    if isinstance(block, TextBlock):
        if block.is_header or block.is_footer:
            return
        _render_text_in_cell(cell, block)
    elif isinstance(block, TableBlock):
        _render_table_in_cell(cell, block, scale)
    elif isinstance(block, ImageBlock):
        _render_image_in_cell(cell, block, page_width, scale)
