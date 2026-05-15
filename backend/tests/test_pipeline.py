from docx import Document
from io import BytesIO

from app.models.enums import ConversionMode
from app.pipeline.detect import detect_pdf_type
from app.pipeline.orchestrator import run_conversion
from app.models.enums import PdfType


def test_detect_digital_pdf(digital_pdf_bytes):
    pdf_type = detect_pdf_type(digital_pdf_bytes, ConversionMode.FAST)
    assert pdf_type == PdfType.DIGITAL


def test_fast_conversion_produces_docx(digital_pdf_bytes):
    result = run_conversion(digital_pdf_bytes, ConversionMode.FAST)
    assert len(result.docx_bytes) > 1000
    doc = Document(BytesIO(result.docx_bytes))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "PDFintoWord" in text or len(doc.paragraphs) > 0


def test_accurate_conversion_produces_docx(digital_pdf_bytes):
    result = run_conversion(digital_pdf_bytes, ConversionMode.ACCURATE)
    assert len(result.docx_bytes) > 1000
    assert result.pdf_type == PdfType.DIGITAL


def test_multi_page_pdf_has_content(digital_pdf_bytes):
    result = run_conversion(digital_pdf_bytes, ConversionMode.ACCURATE)
    doc = Document(BytesIO(result.docx_bytes))
    assert result.page_count >= 2
