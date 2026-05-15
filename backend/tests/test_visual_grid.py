from io import BytesIO

from docx import Document

from app.models.enums import ConversionMode
from app.pipeline.models import DocumentModel
from app.pipeline.orchestrator import run_conversion
from app.pipeline.visual.pipeline import run_visual_reconstruction


def test_accurate_uses_layout_tables(digital_pdf_bytes):
    result = run_conversion(digital_pdf_bytes, ConversionMode.ACCURATE)
    assert len(result.docx_bytes) > 2000
    doc = Document(BytesIO(result.docx_bytes))
    assert len(doc.tables) >= 1 or len(doc.paragraphs) > 0


def test_visual_reconstruction_sets_grid(digital_pdf_bytes):
    from app.pipeline.layout.extractor import extract_layout

    pages = extract_layout(digital_pdf_bytes, ConversionMode.ACCURATE)
    doc = DocumentModel(pages=pages, pdf_type="digital")
    doc = run_visual_reconstruction(doc, digital_pdf_bytes, ConversionMode.ACCURATE)
    assert doc.use_visual_grid
    assert doc.pages[0].grid is not None
    assert doc.pages[0].grid.row_count >= 1
    assert doc.pages[0].canvas is not None
